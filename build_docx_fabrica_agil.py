from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path


OUT = Path(__file__).resolve().parent / "Fabrica-Agil-MVP.docx"


sections = [
    {
        "title": "1. Tese do Negócio",
        "blocks": [
            ("h3", "Fato"),
            ("p", "Pequenas e médias indústrias sofrem com perdas operacionais recorrentes: atraso de pedidos, retrabalho, desperdício de material, baixa previsibilidade, gargalos produtivos, setups demorados, falta de padrão e ausência de rotina de melhoria."),
            ("p", "Essas dores normalmente não acontecem por falta de teoria disponível. Métodos como Lean Manufacturing, Sistema Toyota de Produção, 5S, Kaizen, Teoria das Restrições e mapeamento de fluxo de valor já existem e são amplamente conhecidos no ambiente industrial."),
            ("h3", "Inferência"),
            ("p", "O problema das PMEs industriais não é acesso a conceitos. O problema é transformar esses conceitos em decisões simples, priorizadas e aplicáveis na rotina da fábrica."),
            ("p", "O dono ou gestor sabe que a fábrica está travada, mas normalmente não sabe responder com clareza:"),
            ("ul", ["Qual é o gargalo principal agora?", "Qual desperdício está custando mais dinheiro?", "O que deve ser corrigido primeiro?", "Como orientar a equipe sem contratar uma consultoria cara?", "Como acompanhar se a melhoria funcionou?"]),
            ("h3", "Recomendação"),
            ("p", "O Fábrica Ágil deve nascer como um agente de IA especialista em produtividade industrial para PMEs."),
            ("quote", "Fábrica Ágil é um consultor industrial de IA para pequenas e médias fábricas. Ele diagnostica gargalos, identifica desperdícios e transforma boas práticas de produção em planos de ação simples, rápidos e executáveis."),
            ("p", "O produto não deve começar como ERP, MES, BI, curso tradicional ou consultoria convencional. O foco inicial deve ser diagnosticar o problema operacional, priorizar a ação e guiar a execução semanal."),
        ],
    },
    {
        "title": "2. Problema e Dor",
        "blocks": [
            ("h3", "Problema Central"),
            ("p", "Pequenas e médias indústrias perdem produtividade e margem porque não conseguem identificar, priorizar e resolver os problemas básicos de gestão da produção."),
            ("h3", "Dor Principal"),
            ("p", "O gestor da fábrica precisa produzir mais, atrasar menos e desperdiçar menos, mas não tem tempo, dinheiro ou conhecimento técnico suficiente para contratar e implementar uma consultoria industrial completa."),
            ("h3", "Dores Específicas"),
            ("ol", ["Atraso na entrega de pedidos.", "Retrabalho frequente.", "Desperdício de matéria-prima.", "Excesso de estoque ou falta de material no momento errado.", "Gargalos invisíveis ou mal diagnosticados.", "Setup demorado entre pedidos, produtos ou lotes.", "Falta de padrão operacional.", "Baixa produtividade por funcionário.", "Dono centralizando decisões operacionais.", "Equipe sem clareza de prioridade.", "Falta de indicadores simples.", "Gestor sabe que há perda, mas não sabe medir.", "Medo de consultoria cara e complexa.", "Dificuldade de aplicar Lean, 5S, Kaizen e gestão de gargalos na prática."]),
            ("h3", "Sintoma Visível"),
            ("p", "O dono diz frases como:"),
            ("ul", ["Minha produção vive atrasada.", "Tenho funcionário, mas a fábrica não rende.", "A equipe trabalha o dia inteiro e mesmo assim não entrega.", "Perco muito material e não sei onde.", "Já tentei organizar, mas volta tudo à bagunça.", "Consultoria é caro demais para minha realidade.", "Eu precisava de alguém me dizendo o que fazer primeiro."]),
            ("h3", "Causa Raiz"),
            ("p", "O problema não é apenas falta de ferramenta. A causa raiz é a combinação de:"),
            ("ul", ["baixa maturidade de gestão industrial;", "ausência de diagnóstico simples;", "falta de rotina de acompanhamento;", "conhecimento técnico concentrado fora da empresa;", "dificuldade de traduzir método em ação;", "pouca disciplina de medição;", "sobrecarga do dono ou gestor."]),
            ("h3", "Contra-argumento Forte"),
            ("p", "O risco do negócio é acreditar que a PME industrial quer mais um sistema. Muitas fábricas já têm WhatsApp, planilha, ERP simples, caderno, controles improvisados e promessas anteriores de melhoria. Se o Fábrica Ágil for percebido como mais uma tela para preencher dados, a adesão será baixa."),
            ("h3", "Implicação"),
            ("quote", "Responda algumas perguntas e receba um plano prático para melhorar sua fábrica esta semana."),
        ],
    },
    {
        "title": "3. ICP - Cliente Ideal",
        "blocks": [
            ("h3", "ICP Inicial Recomendado"),
            ("p", "Micro e pequena indústria de produção discreta, com processos visíveis, gargalos recorrentes e baixa maturidade de gestão."),
            ("h3", "Perfil da Empresa"),
            ("ul", ["Setores iniciais: moveleiro, marcenaria industrial, madeira, metalmecânica leve, serralheria produtiva e pequenas fábricas sob encomenda.", "Tamanho: 8 a 80 funcionários.", "Estrutura: dono ou gestor ainda muito presente na operação.", "Maturidade: usa planilhas, caderno, WhatsApp ou ERP básico.", "Processo: tem etapas produtivas claras, como corte, montagem, acabamento, embalagem, expedição, solda, pintura, usinagem ou montagem.", "Dor: atraso, retrabalho, gargalo, desperdício ou baixa previsibilidade.", "Capacidade de pagamento: consegue pagar assinatura se enxergar retorno rápido."]),
            ("h3", "Persona Principal: Dono-operador industrial"),
            ("ul", ["Tem conhecimento prático da fábrica, mas pouca metodologia estruturada.", "Está sobrecarregado.", "Decide compras, equipe, produção, entrega e cliente.", "Quer resultado prático, não teoria.", "Tem resistência a projetos longos.", "Prefere linguagem simples e direta."]),
            ("quote", "Eu sei que dá para produzir melhor, mas não sei por onde começar e não tenho tempo para parar a fábrica."),
            ("h3", "Persona Secundária: Gestor de produção ou encarregado"),
            ("ul", ["Cuida do dia a dia do chão de fábrica.", "Recebe pressão por prazo.", "Conhece os problemas, mas não tem autoridade ou método para priorizar.", "Precisa de ferramenta para justificar mudanças ao dono."]),
            ("quote", "Se eu conseguir mostrar o gargalo com clareza, fica mais fácil mudar a rotina."),
            ("h3", "Melhor Nicho Para Entrada"),
            ("p", "Indústrias moveleiras e marcenarias produtivas sob encomenda ou semi-seriadas."),
            ("ul", ["Processo produtivo fácil de mapear.", "Gargalos comuns e visíveis.", "Alto impacto de retrabalho, atraso e desperdício.", "Dor forte em prazo e organização.", "Linguagem do problema é familiar para os fundadores.", "Permite validar antes de expandir para outros setores."]),
            ("h3", "Anti-ICP"),
            ("ul", ["indústria grande com ERP/MES robusto;", "empresa muito pequena sem recorrência produtiva;", "negócio artesanal sem processo repetível;", "indústria de processo contínuo complexo;", "empresa que não aceita registrar informações básicas;", "cliente que quer uma consultoria totalmente personalizada desde o primeiro dia."]),
        ],
    },
    {
        "title": "4. Ideação do Negócio",
        "blocks": [
            ("h3", "Direção Principal"),
            ("p", "O Fábrica Ágil deve combinar três camadas:"),
            ("ol", ["Agente de IA industrial.", "Método educacional prático.", "Rotina de execução e acompanhamento."]),
            ("h3", "Opção 1 - Agente de IA para Produtividade Industrial"),
            ("p", "Um chat/agente especializado que conversa com o gestor, faz diagnóstico, identifica gargalos, explica problemas e gera planos de ação."),
            ("ul", ["Valor: atendimento imediato, baixo custo comparado à consultoria, escalabilidade e sensação de suporte técnico contínuo.", "Risco: se for genérico, vira apenas ChatGPT para fábrica; se depender de dados demais, o usuário abandona.", "Requisito: o agente precisa ter método próprio, perguntas guiadas e saídas padronizadas."]),
            ("h3", "Opção 2 - Educação Estilo G4 para Industriais"),
            ("p", "Uma trilha prática para ensinar donos de pequenas fábricas a aplicar produtividade industrial sem linguagem acadêmica."),
            ("ul", ["Valor: cria autoridade, facilita venda inicial e pode gerar caixa antes do SaaS completo.", "Risco: se virar só curso, perde diferenciação; conteúdo sem acompanhamento não garante implementação.", "Requisito: educação deve estar conectada ao plano de ação. Cada módulo precisa levar a uma tarefa prática."]),
            ("h3", "Opção 3 - Consultoria Industrial Produtizada"),
            ("p", "Pacotes de diagnóstico, mentoria e auditoria com apoio da plataforma."),
            ("ul", ["Valor: aumenta ticket médio, gera aprendizado rápido com clientes reais e cria casos de sucesso.", "Risco: pode prender os fundadores em entrega manual e escala menos que software.", "Requisito: toda entrega manual deve virar regra, template ou fluxo dentro do produto."]),
            ("h3", "Opção 4 - Software Completo de Gestão Industrial"),
            ("p", "Sistema com dashboards, indicadores, apontamento, PCP, integrações, ERP, estoque e controle operacional."),
            ("ul", ["Valor: mercado grande e possibilidade de virar plataforma robusta no longo prazo.", "Risco: alto custo de desenvolvimento, venda mais longa, concorrência com ERP/MES e desvio do problema inicial.", "Requisito: não deve ser o MVP. Pode ser expansão futura."]),
            ("h3", "Recomendação de Ideação"),
            ("quote", "Começar com agente de IA + diagnóstico guiado + educação aplicada + plano de ação semanal. Não começar com ERP, MES, BI complexo, integrações, aplicativo de apontamento ou automação operacional completa."),
        ],
    },
    {
        "title": "5. Produto - Conceito do MVP",
        "blocks": [
            ("h3", "Nome do MVP"),
            ("p", "Fábrica Ágil Sprint."),
            ("h3", "Promessa"),
            ("quote", "Em 30 dias, o Fábrica Ágil ajuda sua fábrica a identificar o principal gargalo, reduzir desperdícios e executar um plano simples de melhoria da produção."),
            ("h3", "Objetivo do MVP"),
            ("p", "Validar se PMEs industriais pagam por um agente de IA que diagnostica problemas produtivos e gera planos de ação práticos."),
            ("h3", "Hipótese Principal"),
            ("p", "Se o gestor industrial responder um diagnóstico simples e receber um plano de ação claro, ele percebe valor suficiente para pagar uma assinatura mensal."),
            ("h3", "Hipóteses Secundárias"),
            ("ol", ["O usuário aceita conversar com IA sobre problemas da fábrica.", "O usuário consegue responder perguntas básicas sem precisar integrar sistemas.", "O plano de ação gerado é útil mesmo sem dados perfeitos.", "A educação aumenta adesão ao plano.", "Acompanhamento semanal aumenta retenção."]),
        ],
    },
    {
        "title": "6. Escopo do MVP",
        "blocks": [
            ("h3", "Funcionalidades Obrigatórias"),
            ("p", "1. Onboarding da Fábrica: coletar nome da empresa, setor, quantidade de funcionários, tipo de produção, principais etapas produtivas, maior dor atual, volume aproximado de produção, prazo médio de entrega e principais problemas percebidos."),
            ("p", "2. Diagnóstico Guiado: perguntas por fluxo de produção, gargalos, qualidade e retrabalho, desperdício, estoque e materiais, setup, pessoas e rotina, indicadores, entrega e prazo."),
            ("p", "3. Agente de IA Industrial: responder dúvidas, explicar conceitos em linguagem simples, aprofundar perguntas, gerar plano de ação, revisar progresso semanal e sugerir próxima melhoria."),
            ("p", "4. Plano de Ação: problema identificado, causa provável, ação recomendada, responsável sugerido, prazo, dificuldade, impacto esperado, indicador e passo a passo."),
            ("p", "5. Trilha Educacional Aplicada: módulos curtos conectados a tarefas práticas."),
            ("p", "6. Acompanhamento Semanal: atualização de status, resultado observado, avaliação da IA, ajuste do plano e nova prioridade da semana."),
            ("h3", "Módulos Educacionais Iniciais"),
            ("ol", ["Como encontrar o gargalo da fábrica.", "Os 8 desperdícios na pequena indústria.", "Como organizar fluxo sem software caro.", "Como reduzir retrabalho.", "Como diminuir setup.", "Como fazer uma reunião diária de produção.", "Como acompanhar 3 indicadores simples.", "Como rodar uma melhoria Kaizen por semana."]),
            ("h3", "Funcionalidades Fora do MVP"),
            ("ul", ["integração com ERP;", "apontamento de produção por operador;", "controle de estoque completo;", "módulo financeiro;", "BI avançado;", "aplicativo mobile nativo;", "IoT ou leitura de máquina;", "automação com WhatsApp;", "marketplace de consultores;", "comunidade;", "multiempresa;", "permissão avançada por usuário;", "integrações estilo G4OS."]),
        ],
    },
    {
        "title": "7. Jornada do Usuário",
        "blocks": [
            ("h3", "Primeiro Acesso"),
            ("ol", ["Usuário cria conta.", "Informa dados básicos da fábrica.", "Escolhe sua maior dor: atraso, retrabalho, desperdício, gargalo, baixa produtividade ou desorganização.", "Responde diagnóstico guiado.", "Recebe radar da fábrica.", "Conversa com agente de IA.", "Recebe plano de ação de 7 dias."]),
            ("h3", "Semana 1"),
            ("ol", ["Usuário executa primeira ação.", "Marca status.", "Informa evidência simples.", "IA pergunta o que mudou.", "Sistema ajusta prioridade."]),
            ("h3", "Dia 30"),
            ("ol", ["Sistema mostra progresso.", "Compara score inicial e atual.", "Mostra ações concluídas.", "Sugere continuidade.", "Oferece assinatura mensal ou plano com mentoria."]),
        ],
    },
    {
        "title": "8. Modelo de Receita",
        "blocks": [
            ("h3", "Plano Inicial"),
            ("p", "Assinatura mensal para acesso ao agente, diagnóstico, plano de ação e trilha prática."),
            ("ul", ["Básico: R$ 197 a R$ 297 por mês.", "Pro: R$ 497 a R$ 797 por mês.", "Premium com mentoria: R$ 1.500 a R$ 3.000 por mês ou pacote fechado."]),
            ("h3", "Oferta Beta"),
            ("quote", "Fábrica Ágil Sprint 30 dias - diagnóstico da fábrica, plano de ação por IA e acompanhamento semanal."),
            ("ul", ["R$ 497 a R$ 997 sem acompanhamento individual pesado.", "R$ 1.500 a R$ 3.000 com diagnóstico individual e reunião técnica."]),
            ("h3", "Upsell"),
            ("ul", ["mentoria mensal;", "auditoria online;", "diagnóstico aprofundado;", "pacote de implantação de rotina de produção;", "treinamento para líderes."]),
        ],
    },
    {
        "title": "9. Diferenciação",
        "blocks": [
            ("h3", "Diferença Para Consultoria"),
            ("p", "Consultoria tradicional é cara, demorada, depende de agenda e muitas vezes entrega relatório complexo. O Fábrica Ágil é acessível, imediato, recorrente, prático e orientado a plano semanal."),
            ("h3", "Diferença Para ERP/MES"),
            ("p", "ERP/MES registra operação, controla dados e exige implantação. O Fábrica Ágil interpreta problemas, recomenda ações, educa o gestor, prioriza melhorias e não exige integração inicial."),
            ("h3", "Diferença Para Curso"),
            ("p", "Curso ensina conteúdo. O Fábrica Ágil diagnostica a fábrica, recomenda o que fazer, acompanha execução e ensina apenas o necessário para agir."),
        ],
    },
    {
        "title": "10. Roadmap",
        "blocks": [
            ("h3", "Fase 0 - Validação Manual"),
            ("p", "Prazo: 2 a 3 semanas. Objetivo: validar dor, linguagem e disposição de pagamento."),
            ("ul", ["roteiro de entrevista;", "diagnóstico em formulário;", "plano de ação gerado com apoio manual;", "5 a 10 conversas com indústrias;", "3 clientes beta pagos."]),
            ("h3", "Fase 1 - MVP Funcional"),
            ("p", "Prazo: 4 a 6 semanas. Objetivo: construir versão simples com agente, diagnóstico e plano de ação."),
            ("ul", ["landing page objetiva;", "cadastro/login;", "onboarding;", "diagnóstico guiado;", "chat com agente;", "gerador de plano de ação;", "painel simples de tarefas;", "trilha educacional básica."]),
            ("h3", "Fase 2 - Beta Pago"),
            ("p", "Prazo: 30 a 60 dias. Objetivo: medir uso real e resultado percebido."),
            ("ul", ["5 a 15 indústrias usando;", "acompanhamento semanal;", "coleta de feedback;", "refinamento do diagnóstico;", "biblioteca de ações recomendadas;", "primeiros casos de sucesso."]),
            ("h3", "Fase 3 - SaaS Inicial"),
            ("p", "Objetivo: transformar aprendizado em produto recorrente."),
            ("ul", ["assinatura;", "histórico de diagnósticos;", "score evolutivo;", "templates por setor;", "recomendações melhores por segmento;", "relatórios de progresso."]),
        ],
    },
    {
        "title": "11. Métricas de Validação",
        "blocks": [
            ("h3", "Métricas de Mercado"),
            ("ul", ["quantidade de entrevistas realizadas;", "percentual que reconhece a dor;", "percentual que aceita pagar;", "ticket aceito;", "objeções mais comuns."]),
            ("h3", "Métricas de Produto"),
            ("ul", ["taxa de conclusão do diagnóstico;", "número de planos gerados;", "número de ações criadas;", "número de ações concluídas;", "frequência de uso semanal;", "perguntas feitas ao agente;", "retorno ao produto na segunda semana."]),
            ("h3", "Métricas de Resultado"),
            ("ul", ["redução percebida de atraso;", "redução de retrabalho;", "redução de tempo parado;", "melhoria de organização;", "clareza sobre gargalo;", "economia estimada;", "aumento de produtividade percebido."]),
            ("h3", "Critério de Sucesso do MVP"),
            ("ul", ["pelo menos 5 indústrias pagarem pelo beta;", "pelo menos 70% completarem o diagnóstico;", "pelo menos 50% executarem 2 ou mais ações;", "pelo menos 3 clientes relatarem melhoria operacional concreta;", "pelo menos 2 clientes aceitarem continuar em assinatura."]),
        ],
    },
    {
        "title": "12. Riscos e Alertas",
        "blocks": [
            ("h3", "Risco 1 - Produto Virar Chatbot Genérico"),
            ("p", "Alerta: se o agente responder de forma ampla demais, sem método e sem plano de ação, o valor percebido cai. Mitigação: criar diagnóstico estruturado, biblioteca de problemas e formato padrão de plano."),
            ("h3", "Risco 2 - Exigir Dados Demais"),
            ("p", "Alerta: PMEs industriais podem abandonar se precisarem preencher muitas informações antes de ver valor. Mitigação: começar com diagnóstico leve e perguntas progressivas."),
            ("h3", "Risco 3 - Virar Consultoria Manual"),
            ("p", "Alerta: se toda entrega depender do Lucas, o negócio não escala. Mitigação: documentar cada recomendação e transformar em regra, template ou fluxo do agente."),
            ("h3", "Risco 4 - Venda Muito Conceitual"),
            ("p", "Alerta: Lean Manufacturing com IA pode parecer abstrato para o cliente. Mitigação: vender dor concreta: atrasar menos, reduzir retrabalho, encontrar gargalo, produzir mais com a mesma equipe e organizar a fábrica em 30 dias."),
            ("h3", "Risco 5 - Construir Grande Demais"),
            ("p", "Alerta: tentar copiar o G4OS inteiro aumenta custo e atraso. Mitigação: copiar a lógica de agente e sistema operacional de decisão, não a complexidade de integrações."),
        ],
    },
    {
        "title": "13. Decisões Recomendadas",
        "blocks": [
            ("h3", "Opção A - Começar Como Curso"),
            ("p", "Trade-off: mais fácil de vender e entregar no curto prazo, mas menor diferenciação tecnológica no longo prazo."),
            ("h3", "Opção B - Começar Como SaaS Completo"),
            ("p", "Trade-off: maior potencial de escala, mas maior risco, custo e tempo até validar."),
            ("h3", "Opção C - Começar Como Agente + Diagnóstico + Sprint"),
            ("p", "Trade-off: menos completo que um sistema industrial tradicional, mas mais rápido para validar, vender e aprender com clientes reais."),
            ("h3", "Recomendação Final"),
            ("p", "Começar pela Opção C. O cliente precisa de clareza e ação, não de complexidade. O MVP deve entregar uma experiência parecida com uma consultoria enxuta, mas operada por IA e sustentada por uma metodologia própria."),
            ("quote", "O Fábrica Ágil será um agente de IA para produtividade industrial de PMEs, com diagnóstico guiado, educação aplicada e plano de ação semanal."),
        ],
    },
    {
        "title": "14. Próxima Entrega Necessária",
        "blocks": [
            ("p", "Para transformar este documento em execução, as próximas entregas devem ser:"),
            ("ol", ["Roteiro de entrevista com donos de fábrica.", "Diagnóstico inicial com perguntas e pontuação.", "Estrutura do prompt/método do agente.", "Mapa de telas do MVP.", "Backlog técnico priorizado.", "Landing page de validação.", "Oferta beta de 30 dias."]),
        ],
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(11, 37, 69)
    title.paragraph_format.space_after = Pt(10)

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    quote = styles.add_style("Strategic Quote", WD_STYLE_TYPE.PARAGRAPH)
    quote.base_style = styles["Normal"]
    quote.font.italic = True
    quote.font.color.rgb = RGBColor(31, 77, 120)
    quote.paragraph_format.left_indent = Inches(0.25)
    quote.paragraph_format.right_indent = Inches(0.15)
    quote.paragraph_format.space_before = Pt(4)
    quote.paragraph_format.space_after = Pt(8)


def add_metadata_table(doc):
    table = doc.add_table(rows=4, cols=2)
    table.autofit = False
    table.style = "Table Grid"
    widths = [Inches(1.55), Inches(4.8)]
    rows = [
        ("Projeto", "Fábrica Ágil"),
        ("Documento", "Problema, ICP, Ideação e MVP"),
        ("Versão", "Base estratégica para edição"),
        ("Data", "29/07/2026"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        for idx, text in enumerate([label, value]):
            cell = row.cells[idx]
            cell.width = widths[idx]
            set_cell_margins(cell)
            if idx == 0:
                set_cell_shading(cell, "F2F4F7")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True
            cell.text = text
            if idx == 0:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True
    doc.add_paragraph()


def add_block(doc, kind, content):
    if kind == "h3":
        doc.add_paragraph(content, style="Heading 3")
    elif kind == "p":
        doc.add_paragraph(content, style="Normal")
    elif kind == "quote":
        doc.add_paragraph(content, style="Strategic Quote")
    elif kind == "ul":
        for item in content:
            doc.add_paragraph(item, style="List Bullet")
    elif kind == "ol":
        for item in content:
            doc.add_paragraph(item, style="List Number")


def build():
    doc = Document()
    style_doc(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Fábrica Ágil - Documento de Problema, ICP, Ideação e MVP")

    subtitle = doc.add_paragraph()
    subtitle.add_run("Base estratégica para construir o MVP do agente de IA industrial para PMEs.").italic = True
    add_metadata_table(doc)

    doc.add_paragraph("Síntese executiva", style="Heading 1")
    doc.add_paragraph("O Fábrica Ágil deve nascer como um agente de IA especialista em produtividade industrial para pequenas e médias empresas. O produto precisa diagnosticar gargalos, identificar desperdícios e transformar boas práticas de produção em planos de ação simples, rápidos e executáveis.", style="Normal")
    doc.add_paragraph("A recomendação central é evitar ERP, MES, BI complexo e integrações pesadas no MVP. O foco inicial deve ser diagnóstico guiado, agente de IA, trilha educacional aplicada e acompanhamento semanal de ações.", style="Normal")

    for section in sections:
        doc.add_paragraph(section["title"], style="Heading 1")
        for kind, content in section["blocks"]:
            add_block(doc, kind, content)

    doc.add_paragraph("Fim do documento", style="Heading 2")
    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
